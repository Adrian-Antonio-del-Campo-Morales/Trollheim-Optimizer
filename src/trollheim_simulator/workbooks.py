"""Persistencia de candidatos, enemigos y resultados de simulación en Excel."""

from __future__ import annotations

from datetime import datetime
import json
import math
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .candidate_catalog import (
    GENERAL_SKILL_DESCRIPTIONS,
    armour_descriptions,
    load_bands,
    weapon_descriptions,
)
from .rules import HOUSE_RULES, POISON_DESCRIPTIONS, PREPARATION_DESCRIPTIONS


FORMAT_VERSION = 3
FORMAT_MARKER = "TROLLHEIM_WORKBOOK_V3"
DATA_SHEET = "_Trollheim"
SUMMARY_SHEET = "Candidato"
ENEMIES_SHEET = "Enemigos"
RESULTS_INDEX_SHEET = "Índice de resultados"
RESULT_SHEET_PREFIX = "Resultado · "

NAVY = "17243A"
TEAL = "287D7A"
GOLD = "E9C46A"
PALE = "EEF4F3"
WHITE = "FFFFFF"
GREY = "667085"
THIN_GREY = Side(style="thin", color="D0D5DD")


class CandidateWorkbookError(ValueError):
    pass


def _configured_equipment(config: dict) -> list[str]:
    result = []
    if config.get("has_helmet"):
        result.append("Casco")
    if config.get("has_luck_amulet"):
        result.append("Amuleto de la suerte")
    if config.get("has_sea_dragon_cloak"):
        result.append("Capa de Dragón Marino")
    result.extend(config.get("preparations", ()))
    for poison in (
        config.get("main_poison", "Sin veneno"),
        config.get("offhand_poison", "Sin veneno"),
    ):
        if poison != "Sin veneno" and poison not in result:
            result.append(poison)
    return result


def _section(ws, row: int, title: str, end_column: int = 8) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_column)
    cell = ws.cell(row, 1, title)
    cell.fill = PatternFill("solid", fgColor=TEAL)
    cell.font = Font(color=WHITE, bold=True, size=11)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 22
    return row + 1


def _label_value(ws, row: int, label: str, value, column: int = 1) -> None:
    label_cell = ws.cell(row, column, label)
    value_cell = ws.cell(row, column + 1, "—" if value in (None, "") else value)
    label_cell.font = Font(bold=True, color=NAVY)
    label_cell.fill = PatternFill("solid", fgColor=PALE)
    label_cell.border = value_cell.border = Border(bottom=THIN_GREY)
    label_cell.alignment = Alignment(wrap_text=True, vertical="top")
    value_cell.alignment = Alignment(wrap_text=True, vertical="top")
    text = "—" if value in (None, "") else str(value)
    approximate_width = 28 if column >= 5 else 34
    value_lines = sum(max(1, math.ceil(len(part) / approximate_width)) for part in text.splitlines())
    label_width = 20 if column == 1 else 17
    label_lines = max(1, math.ceil(len(label) / label_width))
    lines = max(value_lines, label_lines)
    ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 15, 15 * lines)


def _attributes(ws, row: int, config: dict) -> int:
    """Dibuja la banda de atributos con seis casillas idénticas."""
    stats = ("HA", "F", "R", "H", "I", "A")
    for column, stat in enumerate(stats, 1):
        header = ws.cell(row, column, stat)
        header.fill = PatternFill("solid", fgColor=GOLD)
        header.font = Font(bold=True, color=NAVY)
        header.alignment = Alignment(horizontal="center", vertical="center")
        value = ws.cell(row + 1, column, config.get(stat, 0))
        value.font = Font(bold=True, size=12)
        value.alignment = Alignment(horizontal="center", vertical="center")
        value.border = Border(bottom=THIN_GREY)
    ws.row_dimensions[row].height = 20
    ws.row_dimensions[row + 1].height = 22
    return row + 3


def _build_summary(workbook, payload: dict) -> None:
    ws = workbook.create_sheet(SUMMARY_SHEET, 0)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    ws.merge_cells("A1:H2")
    title = ws["A1"]
    title.value = "TROLLHEIM · FICHA DE SIMULACIÓN"
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.font = Font(color=WHITE, bold=True, size=18)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 25
    ws["A3"] = f"Actualizado: {datetime.now():%d/%m/%Y %H:%M}"
    ws["A3"].font = Font(italic=True, color=GREY)

    config = payload["config"]
    metadata = payload.get("candidate", {})
    row = _section(ws, 5, "CANDIDATO")
    _label_value(ws, row, "Nombre", metadata.get("name", "Candidato"), 1)
    _label_value(ws, row, "Banda", metadata.get("band_name", "Selección libre"), 5)
    row += 1
    _label_value(ws, row, "Guerrero", metadata.get("profile_name", "Perfil libre"), 1)
    _label_value(ws, row, "Tipo", metadata.get("profile_type", "—"), 5)
    row += 2

    selected_house_rules = [
        HOUSE_RULES[key]
        for key, enabled in (payload.get("house_rules") or {}).items()
        if enabled and key in HOUSE_RULES
    ]
    if selected_house_rules:
        row = _section(ws, row, "REGLAS DE LA CASA ACTIVAS")
        for rule in selected_house_rules:
            _label_value(ws, row, rule["name"], rule["description"], 1)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
            row += 1
        row += 1

    row = _section(ws, row, "ATRIBUTOS")
    row = _attributes(ws, row, config)

    row = _section(ws, row, "EQUIPO Y HABILIDADES")
    _label_value(ws, row, "Mano principal", config.get("main_weapon"), 1)
    _label_value(ws, row, "Material", config.get("main_weapon_material"), 5)
    row += 1
    _label_value(ws, row, "Mano secundaria", config.get("off_hand"), 1)
    _label_value(ws, row, "Material", config.get("offhand_material"), 5)
    row += 1
    _label_value(ws, row, "Armadura", config.get("armor"), 1)
    _label_value(ws, row, "Equipamiento", ", ".join(_configured_equipment(config)) or "Ninguno", 5)
    row += 1
    _label_value(ws, row, "Habilidades", ", ".join(config.get("skills", ())) or "Ninguna", 1)
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
    row += 2

    row = _section(ws, row, "CONSULTA RÁPIDA")
    weapon_notes = weapon_descriptions()
    armour_notes = armour_descriptions()
    skill_notes = dict(GENERAL_SKILL_DESCRIPTIONS)
    for band in load_bands():
        skill_notes.update((skill.name, skill.description) for skill in band.skills)
    quick_rows = []
    for label, item, descriptions in (
        ("Mano principal", config.get("main_weapon"), weapon_notes),
        ("Mano secundaria", config.get("off_hand"), weapon_notes),
        ("Armadura", config.get("armor"), armour_notes),
    ):
        description = descriptions.get(item, "")
        if item and item not in ("Ninguna", "Sin Armadura") and description:
            quick_rows.append((f"{label}: {item}", description))
    equipment_notes = {
        **armour_notes,
        **PREPARATION_DESCRIPTIONS,
        **POISON_DESCRIPTIONS,
    }
    quick_rows.extend(
        (f"Equipamiento: {item}", equipment_notes.get(item, ""))
        for item in _configured_equipment(config)
        if equipment_notes.get(item)
    )
    quick_rows.extend(
        (f"Habilidad: {skill}", skill_notes.get(skill, "Sin descripción canónica breve."))
        for skill in config.get("skills", ())
    )
    if not quick_rows:
        quick_rows.append(("Equipo y habilidades", "No hay reglas breves adicionales que consultar."))
    for label, description in quick_rows:
        _label_value(ws, row, label, description, 1)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        ws.row_dimensions[row].height = max(
            ws.row_dimensions[row].height or 15,
            15 * max(1, math.ceil(len(description) / 105)),
        )
        row += 1
    row += 1

    source_rules = metadata.get("rules") or ()
    fixed = metadata.get("fixed_equipment") or ()
    restrictions = metadata.get("restrictions") or ()
    if source_rules or fixed or restrictions:
        row = _section(ws, row, "REFERENCIA DEL PERFIL")
        for label, values in (("Equipo fijo", fixed), ("Restricciones", restrictions), ("Reglas", source_rules)):
            _label_value(ws, row, label, "\n".join(f"• {value}" for value in values) or "—", 1)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
            characters = sum(len(str(value)) for value in values)
            wrapped_lines = max(len(values), math.ceil(characters / 105))
            ws.row_dimensions[row].height = max(22, 15 * max(1, wrapped_lines))
            row += 1
        row += 1

    widths = {"A": 18, "B": 18, "C": 18, "D": 18, "E": 18, "F": 18, "G": 13, "H": 13}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.print_title_rows = "1:3"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def _enemy_quick_rows(config):
    notes = weapon_descriptions()
    equipment_notes = {
        **armour_descriptions(),
        **PREPARATION_DESCRIPTIONS,
        **POISON_DESCRIPTIONS,
    }
    skill_notes = dict(GENERAL_SKILL_DESCRIPTIONS)
    for band in load_bands():
        skill_notes.update((skill.name, skill.description) for skill in band.skills)
    rows = []
    for label, key in (("Mano principal", "main_weapon"), ("Mano secundaria", "off_hand")):
        item = config.get(key)
        if item and item != "Ninguna" and notes.get(item):
            rows.append((f"{label}: {item}", notes[item]))
    rows.extend(
        (f"Equipamiento: {item}", equipment_notes.get(item, ""))
        for item in _configured_equipment(config)
        if equipment_notes.get(item)
    )
    rows.extend(
        (f"Habilidad: {skill}", skill_notes.get(skill, "Sin descripción canónica breve."))
        for skill in config.get("skills", ())
    )
    return rows


def _build_enemies(workbook, payload):
    ws = workbook.create_sheet(ENEMIES_SHEET, 1)
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H2")
    ws["A1"] = "TROLLHEIM · ENEMIGOS"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color=WHITE, bold=True, size=18)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    data = payload.get("enemies") or {}
    mode = data.get("mode", "sample")
    row = _section(ws, 4, "CONFIGURACIÓN")
    _label_value(ws, row, "Modo", "Muestra aleatoria" if mode == "sample" else "Perfiles manuales", 1)
    _label_value(ws, row, "Nivel", data.get("level", 0), 5)
    row += 2
    if mode == "sample":
        row = _section(ws, row, "GENERACIÓN ALEATORIA")
        _label_value(ws, row, "Dificultades", ", ".join(data.get("difficulties", ())) or "Ninguna", 1)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        row += 1
        _label_value(
            ws, row, "Funcionamiento",
            "El simulador genera en cada combate un perfil ponderado, equipo legal y las mejoras indicadas por el nivel.", 1,
        )
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
    else:
        profiles = data.get("profiles") or ()
        for index, config in enumerate(profiles, 1):
            row = _section(ws, row, f"ENEMIGO {index} · {config.get('enemy_name', f'Enemigo {index}')}")
            _label_value(ws, row, "Nombre", config.get("enemy_name", f"Enemigo {index}"), 1)
            profile = find_profile_for_workbook(config)
            _label_value(ws, row, "Perfil", profile, 5)
            row += 1
            row = _section(ws, row, "ATRIBUTOS")
            row = _attributes(ws, row, config)
            _label_value(ws, row, "Mano principal", config.get("main_weapon"), 1)
            _label_value(ws, row, "Mano secundaria", config.get("off_hand"), 5)
            row += 1
            _label_value(ws, row, "Armadura", config.get("armor"), 1)
            _label_value(ws, row, "Equipamiento", ", ".join(_configured_equipment(config)) or "Ninguno", 5)
            row += 1
            _label_value(ws, row, "Habilidades", ", ".join(config.get("skills", ())) or "Ninguna", 1)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
            row += 1
            for label, description in _enemy_quick_rows(config):
                _label_value(ws, row, label, description, 1)
                ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
                row += 1
            row += 1
    for column, width in {"A": 18, "B": 18, "C": 18, "D": 18, "E": 18, "F": 18, "G": 13, "H": 13}.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A4"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def find_profile_for_workbook(config):
    from .candidate_catalog import find_profile
    profile = find_profile(
        config.get("enemy_band_id", ""), config.get("enemy_profile_id", "")
    )
    return f"{profile.band_name} · {profile.name}" if profile else "Selección libre"


def _result_sheet_name(title: str, used: set[str]) -> str:
    base = f"{RESULT_SHEET_PREFIX}{title}"[:31]
    name = base
    suffix = 2
    while name in used:
        tail = f" {suffix}"
        name = f"{base[:31 - len(tail)]}{tail}"
        suffix += 1
    used.add(name)
    return name


def _build_result_sheet(workbook, result: dict, sheet_name: str) -> None:
    ws = workbook.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A8"
    headers = tuple(result.get("headers") or ())
    end_column = max(6, len(headers))
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=end_column)
    ws["A1"] = f"TROLLHEIM · {result.get('title', 'RESULTADOS').upper()}"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color=WHITE, bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    _label_value(ws, 4, "Fecha", result.get("generated_at", "—"), 1)
    _label_value(ws, 4, "Iteraciones", result.get("iterations", 0), 5)
    _label_value(ws, 5, "Rival", result.get("opponent", "—"), 1)
    _label_value(ws, 5, "Vista", result.get("view", "—"), 5)
    for column, header in enumerate(headers, 1):
        cell = ws.cell(7, column, header)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_index, values in enumerate(result.get("rows") or (), 8):
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_index, column, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(value, float):
                cell.number_format = "0.00"
    for column in range(1, max(1, len(headers)) + 1):
        letter = get_column_letter(column)
        longest = max(
            [len(str(ws.cell(row, column).value or "")) for row in range(7, ws.max_row + 1)]
            or [12]
        )
        ws.column_dimensions[letter].width = min(42, max(13, longest + 2))
    ws.auto_filter.ref = f"A7:{ws.cell(max(7, ws.max_row), max(1, len(headers))).coordinate}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def _build_results(workbook, results) -> None:
    ws = workbook.create_sheet(RESULTS_INDEX_SHEET)
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F2")
    ws["A1"] = "RESULTADOS GUARDADOS"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(color=WHITE, bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A4"] = "Las simulaciones guardadas aparecerán aquí y en sus propias hojas."
    ws.merge_cells("A4:F4")
    headers = ("Fecha", "Simulación", "Iteraciones", "Rival", "Vista", "Hoja")
    for column, value in enumerate(headers, 1):
        cell = ws.cell(6, column, value)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(color=WHITE, bold=True)
    for column, width in enumerate((20, 24, 14, 28, 16, 24), 1):
        ws.column_dimensions[chr(64 + column)].width = width
    ws.freeze_panes = "A7"
    used = set(workbook.sheetnames)
    for row, result in enumerate(results or (), 7):
        sheet_name = _result_sheet_name(str(result.get("title", "Simulación")), used)
        _build_result_sheet(workbook, result, sheet_name)
        values = (
            result.get("generated_at", "—"), result.get("title", "Simulación"),
            result.get("iterations", 0), result.get("opponent", "—"),
            result.get("view", "—"), sheet_name,
        )
        for column, value in enumerate(values, 1):
            ws.cell(row, column, value)
        ws.cell(row, 6).hyperlink = f"#'{sheet_name}'!A1"
        ws.cell(row, 6).style = "Hyperlink"


def save_candidate_workbook(path, payload: dict) -> Path:
    destination = Path(path)
    workbook = Workbook()
    del workbook["Sheet"]
    payload = {
        **payload,
        "format_version": FORMAT_VERSION,
        "results": list(payload.get("results", ())),
    }
    _build_summary(workbook, payload)
    _build_enemies(workbook, payload)
    _build_results(workbook, payload.get("results", ()))
    data = workbook.create_sheet(DATA_SHEET)
    data.sheet_state = "veryHidden"
    data["A1"] = FORMAT_MARKER
    data["A2"] = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    data["A3"] = FORMAT_VERSION
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination


def load_candidate_workbook(path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=False)
    if DATA_SHEET not in workbook.sheetnames:
        raise CandidateWorkbookError("El libro no contiene una ficha de Trollheim.")
    data = workbook[DATA_SHEET]
    if data["A1"].value != FORMAT_MARKER or data["A3"].value != FORMAT_VERSION:
        raise CandidateWorkbookError("La versión de esta ficha no es compatible.")
    try:
        payload = json.loads(data["A2"].value)
    except (TypeError, json.JSONDecodeError) as exc:
        raise CandidateWorkbookError("Los datos internos de la ficha están dañados.") from exc
    if (
        not isinstance(payload, dict)
        or payload.get("format_version") != FORMAT_VERSION
        or not isinstance(payload.get("config"), dict)
    ):
        raise CandidateWorkbookError("La ficha no contiene un candidato válido.")
    enemies = payload.get("enemies")
    if (
        not isinstance(enemies, dict)
        or enemies.get("mode") not in {"sample", "custom"}
        or not isinstance(enemies.get("profiles"), list)
    ):
        raise CandidateWorkbookError("El libro no contiene una configuración de enemigos válida.")
    results = payload.get("results")
    if not isinstance(results, list) or any(
        not isinstance(result, dict)
        or result.get("format_version") != FORMAT_VERSION
        or result.get("target") not in {"combos", "weapons", "equipment"}
        for result in results
    ):
        raise CandidateWorkbookError("El libro no contiene resultados del formato actual.")
    return payload
