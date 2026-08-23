from pathlib import Path

from openpyxl import load_workbook
import pytest

from trollheim_simulator.candidate_catalog import (
    equipment_costs_for_profile,
    equipment_options_for_profile,
    find_profile,
    load_bands,
)
from trollheim_simulator.ui import TrollheimApp
from trollheim_simulator.rules import SKILLS
from trollheim_simulator.workbooks import (
    CandidateWorkbookError,
    DATA_SHEET,
    ENEMIES_SHEET,
    FORMAT_VERSION,
    RESULTS_INDEX_SHEET,
    RESULT_SHEET_PREFIX,
    SUMMARY_SHEET,
    load_candidate_workbook,
    save_candidate_workbook,
)


def _payload():
    return {
        "config": {
            "candidate_name": "Rata con papeles",
            "candidate_band_id": "trollheim-skaven",
            "candidate_profile_id": "eshin-assassin",
            "HA": 4, "F": 4, "R": 3, "H": 1, "I": 5, "A": 1,
            "skills": ["Golpe Poderoso"], "main_weapon": "Espada",
            "off_hand": "Daga", "main_weapon_material": "Sin material",
            "offhand_material": "Sin material", "armor": "Armadura Ligera",
            "has_helmet": True,
        },
        "candidate": {
            "name": "Rata con papeles", "band_name": "Skaven del Clan Eshin",
            "profile_name": "Asesino", "profile_type": "hero",
            "rules": ["Luchador Consumado: -1 adicional a la salvación."],
        },
        "opponent": {"mode": "Muestra aleatoria", "level": 2, "description": "Baja, Media"},
        "enemies": {
            "mode": "custom", "level": 2, "difficulties": [],
            "profiles": [
                {
                    "enemy_name": "Bruto", "HA": 3, "F": 4, "R": 4,
                    "H": 1, "I": 2, "A": 1, "skills": [],
                    "main_weapon": "Maza", "off_hand": "Ninguna",
                    "main_weapon_material": "Sin material",
                    "offhand_material": "Sin material", "armor": "Sin Armadura",
                    "has_helmet": False,
                },
                {
                    "enemy_name": "Rápido", "HA": 4, "F": 3, "R": 3,
                    "H": 1, "I": 5, "A": 2, "skills": ["Reflejos Felinos"],
                    "main_weapon": "Espada", "off_hand": "Daga",
                    "main_weapon_material": "Sin material",
                    "offhand_material": "Sin material", "armor": "Armadura Ligera",
                    "has_helmet": True,
                },
            ],
        },
    }


def test_equipment_filter_respects_profile_specific_restrictions():
    witch = set(equipment_options_for_profile(
        "lustria-goblins-salvajes", "savage-goblin-witch-doctor"
    ))
    boss = set(equipment_options_for_profile(
        "lustria-goblins-salvajes", "savage-goblin-big-boss"
    ))
    henchman = set(equipment_options_for_profile(
        "lustria-goblins-salvajes", "savage-goblin"
    ))
    assert {"Veneno Negro", "Armadura Kitinoza"} <= witch
    assert {"Amuleto de la suerte", "Loto Negro"} <= boss
    assert "Amuleto de la suerte" not in henchman


def test_general_market_equipment_respects_faction_restrictions():
    matriarch = set(equipment_options_for_profile(
        "trollheim-sisters-of-sigmar", "sigmarite-matriarch"
    ))
    assassin = set(equipment_options_for_profile(
        "trollheim-skaven", "eshin-assassin"
    ))
    assert "Amuleto de la suerte" in matriarch
    assert "Saliva de Araña" in assassin
    assert "Loto Negro" not in matriarch
    assert "Veneno Negro" not in matriarch


def test_equipment_costs_use_band_prices_and_expected_dice_values():
    costs = equipment_costs_for_profile(
        "lustria-elfos-oscuros", "dark-elf-beastmaster"
    )
    assert costs["Armadura Ligera"] == 50.0
    assert costs["Capa de Dragón Marino"] == 57.0
    assert equipment_costs_for_profile()["Amuleto de la suerte"] == 10.0


def test_sea_dragon_cloak_is_special_equipment_not_body_armour():
    corsair = find_profile("lustria-elfos-oscuros", "dark-elf-corsair")
    options = set(equipment_options_for_profile(
        "lustria-elfos-oscuros", "dark-elf-corsair"
    ))
    assert "Capa de Dragón Marino" not in corsair.armors
    assert "Capa de Dragón Marino" in options


def test_canonical_candidate_catalog_is_complete():
    bands = load_bands()
    assert len(bands) == 34
    assert sum(len(band.profiles) for band in bands) == 222
    assassin = find_profile("trollheim-skaven", "eshin-assassin")
    assert assassin.stats == {"HA": 4, "F": 4, "R": 3, "H": 1, "I": 5, "A": 1}
    assert {"Espada", "Daga", "Mayal", "Lanza", "Alabarda"} <= set(assassin.weapons)
    assert "Armadura Ligera" in assassin.armors
    assert assassin.helmet_allowed
    assert "Golpe Poderoso" in assassin.skills


def test_attack_replacement_skills_are_not_hidden_by_starting_attacks():
    mutant = find_profile("trollheim-possessed", "mutant")
    black_knife = find_profile(
        "chaos-streets-deathbringers", "black-knife"
    )
    young_claw = find_profile(
        "chaos-streets-deathbringers", "young-claw"
    )
    assert mutant.stats["A"] == 1
    assert "Estocada Mortal" in mutant.skills
    for profile in (black_knife, young_claw):
        assert profile.stats["A"] == 1
        assert "Golpe Mortal" in profile.skills


def test_natural_profiles_receive_a_neutral_natural_weapon():
    rat_ogre = find_profile("trollheim-skaven", "rat-ogre")
    assert rat_ogre.weapons == ("Arma natural",)
    assert not rat_ogre.armors
    assert not rat_ogre.skills


def test_band_skills_are_contextual_and_profile_access_is_respected():
    bands = load_bands()
    skaven = next(band for band in bands if band.band_id == "trollheim-skaven")
    assert {"Hambre Negra", "Arte del Combate sin Armas"} <= {
        skill.name for skill in skaven.skills
    }
    assassin = find_profile("trollheim-skaven", "eshin-assassin")
    verminkin = find_profile("trollheim-skaven", "verminkin")
    assert "Hambre Negra" in assassin.skills
    assert "Hambre Negra" not in verminkin.skills


def test_every_profile_exposes_exactly_the_skills_from_its_allowed_categories():
    for band in load_bands():
        for profile in band.profiles:
            flattened = tuple(
                skill
                for skills in profile.skills_by_category.values()
                for skill in skills
            )
            assert profile.skills == flattened, (band.name, profile.name)
            assert len(profile.skills) == len(set(profile.skills)), (
                band.name, profile.name,
            )


def test_every_engine_skill_is_available_to_at_least_one_canonical_profile():
    catalog_skills = {
        skill
        for band in load_bands()
        for profile in band.profiles
        for skill in profile.skills
    }
    assert set(SKILLS) <= catalog_skills


def test_high_elf_profiles_keep_distinct_skill_columns_and_restrictions():
    loremaster = find_profile("lustria-altos-elfos", "high-elf-loremaster")
    explorer = find_profile("lustria-altos-elfos", "high-elf-explorer")
    assert "academic" in loremaster.skills_by_category
    assert "shooting" not in loremaster.skills_by_category
    assert "shooting" in explorer.skills_by_category
    assert "academic" not in explorer.skills_by_category
    assert "Maestro de las runas" in loremaster.skills_by_category["special"]
    assert "Maestro de las runas" not in explorer.skills_by_category["special"]
    assert "Suerte" in explorer.skills_by_category["special"]
    assert "Suerte" not in loremaster.skills_by_category["special"]
    assert {
        "Reflejos Felinos", "En Pie de un Salto", "Agilidad élfica",
        "Miniath", "Golpe Infalible", "Suerte",
    } <= set(explorer.skills).intersection(SKILLS)
    assert {
        "Reflejos Felinos", "En Pie de un Salto", "Agilidad élfica",
        "Miniath", "Golpe Infalible",
    } <= set(explorer.skills).intersection(SKILLS)


def test_candidate_workbook_round_trip_uses_only_the_current_format(tmp_path: Path):
    path = tmp_path / "candidato.xlsx"
    payload = _payload()
    payload["config"].update({
        "has_sea_dragon_cloak": True,
        "preparations": ["Sombra Carmesí", "Raíz de Mandrágora"],
        "main_poison": "Loto Negro",
    })
    payload["house_rules"] = {
        "anti_dual": True, "hard_armour": True, "cheap_armour": False,
    }
    save_candidate_workbook(path, payload)
    restored = load_candidate_workbook(path)
    assert restored["config"]["candidate_name"] == "Rata con papeles"
    assert restored["house_rules"]["anti_dual"] is True
    assert [profile["enemy_name"] for profile in restored["enemies"]["profiles"]] == [
        "Bruto", "Rápido"
    ]

    workbook = load_workbook(path)
    assert workbook.sheetnames[:3] == [SUMMARY_SHEET, ENEMIES_SHEET, RESULTS_INDEX_SHEET]
    assert workbook[DATA_SHEET].sheet_state == "veryHidden"
    assert workbook[SUMMARY_SHEET]["A1"].value == "TROLLHEIM · FICHA DE SIMULACIÓN"
    assert workbook[SUMMARY_SHEET].auto_filter.ref is None
    enemy_values = [
        cell.value for row in workbook[ENEMIES_SHEET].iter_rows() for cell in row
    ]
    assert "ENEMIGO 1 · Bruto" in enemy_values
    assert "ENEMIGO 2 · Rápido" in enemy_values
    summary_values = [
        cell.value for row in workbook[SUMMARY_SHEET].iter_rows() for cell in row
    ]
    assert "CONSULTA RÁPIDA" in summary_values
    assert "Habilidad: Golpe Poderoso" in summary_values
    assert "REGLAS DE LA CASA ACTIVAS" in summary_values
    assert "Anti Dos Armas" in summary_values
    assert "Armaduras Duras" in summary_values
    assert "Armaduras Baratas" not in summary_values
    assert any(
        value and all(item in str(value) for item in (
            "Capa de Dragón Marino", "Sombra Carmesí", "Raíz de Mandrágora",
            "Loto Negro",
        ))
        for value in summary_values
    )
    changed = _payload()
    changed["config"]["HA"] = 5
    save_candidate_workbook(path, changed)
    workbook = load_workbook(path, data_only=False)
    assert load_candidate_workbook(path)["config"]["HA"] == 5

    workbook[DATA_SHEET]["A1"] = "TROLLHEIM_WORKBOOK_V2"
    workbook.save(path)
    with pytest.raises(CandidateWorkbookError, match="no es compatible"):
        load_candidate_workbook(path)


def test_workbook_saves_and_loads_each_simulation_on_its_own_sheet(tmp_path: Path):
    path = tmp_path / "simulaciones.xlsx"
    payload = _payload()
    payload["results"] = [
        {
            "format_version": FORMAT_VERSION,
            "target": "combos", "title": "Mejoras",
            "generated_at": "2026-08-23T10:30:00", "iterations": 25000,
            "opponent": "Rival configurable: Bruto", "view": "optimal",
            "headers": ["Mejora", "Mano libre %", "Impacto %"],
            "rows": [["ESTADO BASE", 41.25, 0.0], ["Golpe Poderoso", 49.5, 8.25]],
            "table_data": {"Single": [["ESTADO BASE", 41.25, 0.0], []]},
            "card_data": [{"Single": 41.25}, "Single", {"Single": "Espada"}],
        },
        {
            "format_version": FORMAT_VERSION,
            "target": "weapons", "title": "Configuraciones de armas",
            "generated_at": "2026-08-23T10:35:00", "iterations": 10000,
            "opponent": "Rival configurable: Bruto", "view": "equipment",
            "headers": ["Principal", "Secundaria", "Victoria %"],
            "rows": [["Espada", "Daga", 52.75]],
            "table_data": {"Single": [40.0, []]}, "card_data": None,
        },
    ]

    save_candidate_workbook(path, payload)
    restored = load_candidate_workbook(path)
    assert [result["target"] for result in restored["results"]] == ["combos", "weapons"]

    workbook = load_workbook(path, data_only=False)
    result_sheets = [name for name in workbook.sheetnames if name.startswith(RESULT_SHEET_PREFIX)]
    assert len(result_sheets) == 2
    assert workbook[RESULTS_INDEX_SHEET]["B7"].value == "Mejoras"
    assert workbook[RESULTS_INDEX_SHEET]["C8"].value == 10000
    first_values = [cell.value for row in workbook[result_sheets[0]].iter_rows() for cell in row]
    assert "Golpe Poderoso" in first_values
    assert 49.5 in first_values


def test_result_export_allows_rows_that_do_not_apply_to_every_combat_mode():
    table_data = (
        {
            "Single": [40.0, [["Espada || Ninguna", 45.0, 5.0]]],
            "Shield": [42.0, []],
            "Dual": [43.0, []],
            "TwoHand": [44.0, []],
        },
        {"Single": "Espada", "Shield": "Espada + escudo",
         "Dual": "Espada + daga", "TwoHand": "Arma a dos manos"},
    )

    headers, rows = TrollheimApp._result_export_rows("weapons", table_data)

    weapon_row = next(row for row in rows if row[0] == "Espada")
    assert len(weapon_row) == len(headers)
    assert weapon_row[2] == 45.0
    assert weapon_row[3:6] == (None, None, None)
    assert weapon_row[-2:] == ("Single", "Espada")
